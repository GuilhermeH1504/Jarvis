# Acoes de arquivos: listar pastas, buscar arquivos pelo nome, ler texto de
# arquivos pra o LLM analisar e criar/atualizar planilhas Excel com dados
# fornecidos.

import os
from pathlib import Path

from openpyxl import Workbook, load_workbook

_SEARCH_FOLDERS = [
    "Desktop",
    "Documents",
    "Downloads",
    "Pictures",
    "Videos",
    "Music",
    "OneDrive/Desktop",
    "OneDrive/Documents",
    "OneDrive/Downloads",
]


def open_path(path: str) -> str:
    target = Path(path).expanduser()

    if not target.exists():
        return f"Não encontrei '{path}'."

    os.startfile(target)
    return f"Abrindo {target.name}."


def list_dir(path: str = ".") -> str:
    target = Path(path).expanduser()

    if not target.is_dir():
        return f"'{path}' não é uma pasta válida."

    items = sorted(p.name for p in target.iterdir())
    if not items:
        return f"A pasta '{target}' está vazia."

    return ", ".join(items[:20])


def read_file(path: str, max_chars: int = 4000) -> str:
    """Le o conteudo de um arquivo de texto pra analise. Trunca arquivos
    grandes pra nao estourar o contexto do LLM."""
    target = Path(path).expanduser()

    if not target.exists():
        return f"Não encontrei o arquivo '{path}'."
    if not target.is_file():
        return f"'{path}' não é um arquivo."

    try:
        content = target.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return f"Não consegui ler '{path}': {exc}"

    if len(content) > max_chars:
        content = content[:max_chars] + "\n[...conteúdo truncado...]"

    return content


def find_files(name: str, root: str | None = None, max_results: int = 15) -> str:
    """Procura arquivos cujo nome contenha 'name'. Por padrao busca nas
    pastas comuns do usuario (Area de Trabalho, Documentos, Downloads,
    Imagens, Videos, Musica e as versoes dentro do OneDrive, se existirem).
    Se 'root' for informado (ex: 'D:/Projetos'), busca so ali em vez disso."""
    name_lower = name.lower()
    matches: list[str] = []
    scanned = 0
    scan_cap = 50000

    if root:
        search_bases = [Path(root).expanduser()]
    else:
        candidates = [Path.home() / folder for folder in _SEARCH_FOLDERS]
        seen_resolved = set()
        search_bases = []
        for base in candidates:
            if not base.exists():
                continue
            resolved = base.resolve()
            if resolved in seen_resolved:
                continue
            seen_resolved.add(resolved)
            search_bases.append(base)

    for base in search_bases:
        if not base.exists():
            continue

        for current_dir, _dirs, files in os.walk(base, onerror=lambda _err: None):
            for filename in files:
                scanned += 1
                if name_lower in filename.lower():
                    matches.append(str(Path(current_dir) / filename))
                    if len(matches) >= max_results:
                        break
                if scanned >= scan_cap:
                    break
            if len(matches) >= max_results or scanned >= scan_cap:
                break
        if len(matches) >= max_results:
            break

    if not matches:
        where = root or "Área de Trabalho, Documentos, Downloads, Imagens, Vídeos, Música e OneDrive"
        return f"Nenhum arquivo encontrado com '{name}' em {where}."

    return "\n".join(matches)


def _resolve_excel_path(path: str) -> Path:
    target = Path(path).expanduser()

    if target.parent == Path("."):
        target = Path.home() / "Desktop" / target.name

    if target.suffix.lower() != ".xlsx":
        target = target.with_suffix(".xlsx")

    return target


def write_excel(path: str, rows: list[list], headers: list[str] | None = None) -> str:
    """Cria uma planilha .xlsx com os dados fornecidos, ou - se o arquivo
    indicado ja existir - adiciona as linhas nela em vez de sobrescrever. Se
    o caminho for so um nome de arquivo (sem pasta), usa a Area de Trabalho."""
    target = _resolve_excel_path(path)
    updating_existing = target.exists()

    if updating_existing:
        try:
            workbook = load_workbook(target)
        except OSError as exc:
            return f"Não consegui abrir a planilha existente '{target}': {exc}"
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        if headers:
            sheet.append(headers)

    for row in rows:
        sheet.append(row)

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(target)
    except OSError as exc:
        return f"Não consegui salvar a planilha em '{target}': {exc}"

    verb = "atualizada" if updating_existing else "criada"
    return f"Planilha {verb} em '{target}' ({len(rows)} linha(s) adicionada(s))."
